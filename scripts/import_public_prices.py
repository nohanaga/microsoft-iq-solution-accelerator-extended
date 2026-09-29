import argparse
from datetime import date, datetime, timezone
import hashlib
from html.parser import HTMLParser
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys
from urllib.parse import urljoin, urlparse
from urllib.request import HTTPRedirectHandler, Request, build_opener

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
LANDING = "https://www.worldbank.org/en/research/commodity-markets"
CATALOGUE = "https://datacatalog.worldbank.org/search/dataset/0038238/commodity-prices-history-and-projections"
TERMS = "https://data.worldbank.org/summary-terms-of-use"
HOSTS = {"www.worldbank.org", "thedocs.worldbank.org"}


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8") + b"\n"


def digest(content):
    return hashlib.sha256(content).hexdigest()


def allowed(url):
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.netloc not in HOSTS:
        raise ValueError("Only the official World Bank HTTPS distribution hosts are accepted")
    return url


class OfficialRedirects(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return super().redirect_request(req, fp, code, msg, headers, allowed(newurl))


class DownloadLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = set()

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            href = dict(attrs).get("href", "")
            if urlparse(href).path.endswith("CMO-Historical-Data-Monthly.xlsx"):
                self.links.add(urljoin(LANDING, href))


def download(url):
    request = Request(allowed(url), headers={"User-Agent": "MicroCoffeeDataImporter/1.0"})
    with build_opener(OfficialRedirects()).open(request, timeout=90) as response:
        content = response.read(12 * 1024 * 1024 + 1)
        if len(content) > 12 * 1024 * 1024:
            raise ValueError("Source exceeds the importer download limit")
        return content, allowed(response.geturl())


def normalize(workbook_bytes, first, last):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        if "Monthly Prices" not in workbook.sheetnames:
            raise ValueError("World Bank Monthly Prices worksheet not found; source format changed")
        rows = list(workbook["Monthly Prices"].values)
        selected = {"coffee arabica": "WB_COFFEE_ARABICA", "coffee robusta": "WB_COFFEE_ROBUSTA"}
        columns = {}
        header_index = None
        for index, row in enumerate(rows[:20]):
            found = {re.sub(r"[^a-z]+", " ", str(value).lower()).strip(): position for position, value in enumerate(row)}
            if all(name in found for name in selected):
                columns = {selected[name]: found[name] for name in selected}
                header_index = index
                break
        if header_index is None:
            raise ValueError("Coffee series headings changed; no positional fallback is allowed")
        for column in columns.values():
            units = {str(row[column]).strip().strip("()").replace(" ", "") for row in rows[header_index + 1:header_index + 4]}
            if "$/kg" not in units:
                raise ValueError("Coffee unit is not verified as US dollars per kilogram")
        result = []
        observed = set()
        for row in rows[header_index + 1:]:
            match = re.fullmatch(r"(\d{4})M(\d{2})", str(row[0]).strip())
            if not match:
                continue
            period = date(int(match[1]), int(match[2]), 1)
            if not first <= period <= last:
                continue
            if period in observed:
                raise ValueError("Duplicate source month")
            observed.add(period)
            for series, column in columns.items():
                value = row[column]
                if value is not None and (type(value) not in (int, float) or not math.isfinite(value) or value <= 0):
                    raise ValueError("Unexpected coffee observation; values were not replaced or interpolated")
                result.append({"id": f"{series}_{period:%Y_%m}", "seriesId": series, "period": period.isoformat(),
                               "value": value, "unit": "USD/kg", "dataOrigin": "public-statistic", "sourceUrl": LANDING})
        month = first
        while month <= last:
            if month not in observed:
                raise ValueError(f"Requested month is absent from the source: {month:%Y-%m}")
            month = date(month.year + (month.month == 12), month.month % 12 + 1, 1)
        notes = []
        for sheet in workbook:
            if sheet.title != "Monthly Prices":
                for row in sheet.values:
                    if any("coffee" in str(value).lower() for value in row if value is not None):
                        notes.append({"sheet": sheet.title, "cells": [str(value) for value in row if value is not None]})
        return result, notes
    finally:
        workbook.close()


def acquire(first, last, download_url=None):
    if first > last:
        raise ValueError("from-month must not be after through-month")
    if download_url is None:
        page, _ = download(LANDING)
        parser = DownloadLinks()
        parser.feed(page.decode("utf-8"))
        if len(parser.links) != 1:
            raise ValueError("Monthly workbook link is not unique; specify the official --download-url after reviewing the source page")
        download_url = next(iter(parser.links))
    content, resolved = download(download_url)
    rows, notes = normalize(content, first, last)
    source_hash = digest(content)
    archive = ROOT / "artifacts/downloads" / (source_hash + ".xlsx")
    archive.parent.mkdir(parents=True, exist_ok=True)
    if archive.exists() and digest(archive.read_bytes()) != source_hash:
        raise ValueError("Existing source archive hash mismatch")
    if not archive.exists():
        archive.write_bytes(content)
    output = ROOT / "scenarios/micro-coffee/reference-data/world-bank-coffee" / f"{first:%Y%m}-{last:%Y%m}-{source_hash[:16]}"
    prices = encoded({"rows": rows})
    fields = ["id", "seriesId", "period", "value", "unit", "dataOrigin", "sourceUrl"]
    properties = {field: {"type": "string", "minLength": 1} for field in fields}
    properties["period"] = {"type": "string", "format": "date"}
    properties["value"] = {"type": ["number", "null"], "exclusiveMinimum": 0}
    properties["unit"] = {"type": "string", "const": "USD/kg"}
    properties["dataOrigin"] = {"type": "string", "const": "public-statistic"}
    specification = encoded({"name": "mc_ref_world_bank_coffee_prices", "fields": fields, "primaryKey": ["id"],
                            "dataOrigin": "public-statistic", "schema": {"$schema": "https://json-schema.org/draft/2020-12/schema",
                            "type": "object", "properties": properties, "required": fields, "additionalProperties": False}})
    files = {"prices.json": prices, "table-schema.json": specification}
    if output.exists():
        provenance = json.loads((output / "provenance.json").read_text(encoding="utf-8"))
        if provenance["sourceSha256"] != source_hash or any((output / name).read_bytes() != value for name, value in files.items()):
            raise ValueError("Existing snapshot differs; no implicit overwrite is permitted")
        return output, len(rows)
    provenance = {"schemaVersion": 1, "sourceId": "world-bank-coffee", "dataOrigin": "public-statistic",
                  "title": "World Bank Commodity Price Data (Pink Sheet): Coffee, monthly nominal prices",
                  "publisher": "World Bank", "catalogueUrl": CATALOGUE, "sourceUrl": LANDING,
                  "downloadUrl": download_url, "resolvedDownloadUrl": resolved, "sourceSha256": source_hash,
                  "retrievedAt": datetime.now(timezone.utc).isoformat(), "fromMonth": first.strftime("%Y-%m"),
                  "throughMonth": last.strftime("%Y-%m"), "rows": len(rows), "license": "CC-BY-4.0",
                  "licenseUrl": "https://creativecommons.org/licenses/by/4.0/", "additionalTermsUrl": TERMS,
                  "attribution": "The World Bank: Commodity Price Data (Pink Sheet). Original data providers are identified in the source workbook.",
                  "changes": "Selected monthly Arabica and Robusta nominal-price observations; normalized dates and column names; no interpolation, currency conversion or joining to fictional transactions.",
                  "sourceNotes": notes, "importerSha256": digest(Path(__file__).read_bytes()),
                  "files": {name: digest(value) for name, value in files.items()}, "endorsement": False}
    output.mkdir(parents=True)
    for name, value in {**files, "provenance.json": encoded(provenance)}.items():
        (output / name).write_bytes(value)
    return output, len(rows)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Acquire a versioned World Bank coffee-price snapshot without Fabric access.")
    parser.add_argument("--from-month", default="2020-01")
    parser.add_argument("--through-month", default="2025-12")
    parser.add_argument("--download-url")
    parser.add_argument("--acknowledge-terms", action="store_true", required=True)
    args = parser.parse_args()
    try:
        first = datetime.strptime(args.from_month, "%Y-%m").date()
        last = datetime.strptime(args.through_month, "%Y-%m").date()
        output, count = acquire(first, last, args.download_url)
        print(json.dumps({"reference_path": output.relative_to(ROOT).as_posix(), "rows": count, "dataOrigin": "public-statistic"}, indent=2))
    except Exception as error:
        print(f"Acquisition stopped: {type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(1)